"""Regression coverage for the separate historical clustering worksheet."""

import importlib.util
import tempfile
import unittest
from pathlib import Path

import openpyxl

MODULE_PATH = Path(__file__).parents[1] / 'generate_mhsds_assessment_seed.py'
SPEC = importlib.util.spec_from_file_location('mhsds_seed', MODULE_PATH)
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


class ClusteringDefinitionsTest(unittest.TestCase):
    def test_clustering_only_responses_and_shared_concept_precedence(self):
        book = openpyxl.Workbook()
        main = book.active
        main.title = 'MH Assessment Scales'
        main.append(['Existing tool', 'Existing question', '123456', None, 0, 'Main meaning'])
        cluster = book.create_sheet('Cluster Tools for MH')
        cluster.append(['Shared question', '123456', 0, 'Different clustering meaning'])
        cluster.append([
            'Mental Health Clustering Tool Summary Assessments of Risk and Need rating A score',
            '234567', 0, 'No problem',
        ])
        cluster.append([None, None, 9, 'Unknown'])
        cluster.merge_cells('A2:A3')
        cluster.merge_cells('B2:B3')
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'synthetic_spec.xlsx'
            book.save(path)
            rows = list(MODULE.extract('5.0', path))
        book.close()

        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[0]['assessment_tool_name'], 'Existing tool')
        self.assertEqual(rows[0]['response_description'], 'Main meaning')
        self.assertEqual([r['published_value'] for r in rows[1:]], ['0', '9'])
        self.assertEqual([r['response_description'] for r in rows[1:]], ['No problem', 'Unknown'])
        self.assertEqual(rows[1]['assessment_tool_name'], rows[2]['assessment_tool_name'])
        self.assertEqual([r['source_row'] for r in rows[1:]], [2, 3])


if __name__ == '__main__':
    unittest.main()
